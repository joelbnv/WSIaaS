import csv
import json
from typing import Iterable
from openpyxl import Workbook
import pyarrow as pa
import pyarrow.parquet as pq

from database.databases import MySQLDB, PostgreSQLDB, SQLiteDB
from datetime import date, datetime


def write_batch_csv(records: Iterable, path: str):
    list_records = records
    if not list_records:
        return

    headers = list(list_records[0].__fields__.keys())

    with open(f"{path}/Productos.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f, fieldnames=headers, quotechar='"', lineterminator="\n", quoting=csv.QUOTE_STRINGS,
        )
        writer.writeheader()
        for record in list_records:
            row_dict = record.model_dump(mode="python")
            writer.writerow(row_dict)


def write_batch_jsonl(records: Iterable, path: str):
    with open(f"{path}/Productos.jsonl", "w", encoding="utf-8") as f:
        for record in records:
            f.write(
                json.dumps(record.model_dump(mode="json"), ensure_ascii=False) + "\n"
            )


def write_batch_excel(records: Iterable, path: str):
    list_records = records
    if not list_records:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    first_dict = list_records[0].model_dump(mode="python")
    headers = list(first_dict.keys())
    ws.append(headers)

    for record in list_records:
        row_dict = record.model_dump(mode="json")
        row_values = [row_dict[field] for field in headers]
        ws.append(row_values)

    wb.save(f"{path}/Productos.xlsx")


def write_batch_parquet(records: Iterable, path: str):
    dict_records: list[dict] = [r.model_dump(mode="python") for r in records]
    table = pa.Table.from_pylist(dict_records)
    pq.write_table(table, f"{path}/Productos.parquet")


def write_batch_sqlite(records: Iterable, table: str, db_path: str):
    list_records = records
    if not list_records:
        return

    first_dict = list_records[0].model_dump(mode="python")
    fields = list(first_dict.keys())
    field_list = ", ".join(fields)
    placeholders = ", ".join("?" for _ in fields)

    values = [
        tuple(None if v ==  '' else str(v) for v in rec.model_dump(mode="python").values())
        for rec in list_records
    ]

    with SQLiteDB(db_path) as (conn, cursor):
        cursor.execute(
            f"CREATE TABLE IF NOT EXISTS {table} ({', '.join(f'{f} TEXT' for f in fields)})"
        )

        cursor.execute(f"SELECT 1 FROM {table} LIMIT 1")
        if cursor.fetchone():
            cursor.execute(f"DELETE FROM {table}")
        
        cursor.executemany(
            f"INSERT INTO {table} ({field_list}) VALUES ({placeholders})", values
        )
        conn.commit()


def write_batch_mysql(records: Iterable, table: str, conn_params: dict):
    list_records = records
    if not list_records:
        return

    first_dict = list_records[0].model_dump(mode="python")
    fields = list(first_dict.keys())
    field_list = ", ".join(fields)
    placeholders = ", ".join("%s" for _ in fields)

    values = [
        tuple(None if v ==  '' else str(v) for v in rec.model_dump(mode="python").values())
        for rec in list_records
    ]

    with MySQLDB(conn_params) as (conn, cursor):
        cursor.execute(
            f"CREATE TABLE IF NOT EXISTS {table} ({', '.join(f'{f} TEXT' for f in fields)})"
        )

        cursor.execute(f"SELECT 1 FROM {table} LIMIT 1")
        if cursor.fetchone():
            cursor.execute(f"TRUNCATE TABLE {table}")
        
        cursor.executemany(
            f"INSERT INTO {table} ({field_list}) VALUES ({placeholders})", values
        )
        conn.commit()


def write_batch_postgres(records: Iterable, table: str, conn_params: dict):
    list_records = records
    if not list_records:
        return

    first_dict = list_records[0].model_dump(mode="python")
    fields = list(first_dict.keys())
    field_list = ", ".join(fields)
    placeholders = ", ".join("%s" for _ in fields)

    values = [
        tuple(None if v ==  '' else str(v) for v in rec.model_dump(mode="python").values())
        for rec in list_records
    ]

    with PostgreSQLDB(conn_params) as (conn, cursor):
        cursor.execute(
            f"CREATE TABLE IF NOT EXISTS {table} ({', '.join(f'{f} TEXT' for f in fields)})"
        )

        # Comprobamos si hay datos, y ejecutamos una operación de TRUNCATE para vaciar la tabla si los hay
        cursor.execute(f"SELECT 1 FROM {table} LIMIT 1")
        if cursor.fetchone():
            cursor.execute(f"TRUNCATE TABLE {table}")

        cursor.executemany(
            f"INSERT INTO {table} ({field_list}) VALUES ({placeholders})",
            values
        )
        conn.commit()